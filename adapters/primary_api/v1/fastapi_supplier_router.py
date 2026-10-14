#!/usr/bin/env python3
"""
Module: fastapi_supplier_router.py
Layer: Adapters (Primary API - v1)
Responsibility: REST API endpoint untuk Supplier/Vendor.

PERBAIKAN (2026-10-13):
- Endpoint /export sekarang generate file .xlsx ASLI pakai openpyxl
  ketika format=excel, tidak lagi mengembalikan CSV dengan ekstensi .xlsx
  (yang membuat Excel menolak file dengan "format or file extension is
  not valid").
- Konsisten dengan pola yang sudah dipakai di customer router.

Daftar endpoint:
    /suppliers/suppliers                       CRUD utama
    /suppliers/suppliers/{id}                  Get by ID
    /suppliers/suppliers/{id}/activate|deactivate
    /suppliers/suppliers/{id}/withholding-category
    /suppliers/suppliers/{id}/balance
    /suppliers/suppliers/{id}/history
    /suppliers/suppliers/{id}/attachments      Tab Attachment
    /suppliers/suppliers/{id}/notes            Tab Catatan
    /suppliers/next-code                       Auto-generate code
    /suppliers/stats                           Statistik
    /suppliers/export                          Export CSV/Excel/JSON
    /suppliers/import                          Import
"""

from __future__ import annotations

import csv
import hashlib
import io
import json
import logging
from datetime import date, datetime
from decimal import Decimal
from enum import Enum
from typing import Any
from uuid import UUID

from fastapi import (
    APIRouter,
    Body,
    Depends,
    Header,
    HTTPException,
    Query,
    Request,
    status,
)
from fastapi.responses import Response
from pydantic import BaseModel, ConfigDict, Field

from adapters.dependency_provider import get_service
from adapters.primary_api.common.fastapi_auth_jwt_middleware import (
    TokenPayload,
    get_current_legal_entity,
    get_current_user,
)
from application.service_layer.service_supplier import (
    Supplier,
    SupplierHasTransactionsError,
    SupplierNotFoundError,
    SupplierService,
    SupplierServiceError,
)

logger = logging.getLogger(__name__)


# ============================================================================
# IDEMPOTENCY MANAGER
# ============================================================================


class IdempotencyManager:
    """Simple in-memory idempotency manager untuk FastAPI endpoints."""

    def __init__(self):
        self._storage: dict[str, tuple[str, datetime]] = {}
        self._ttl_seconds = 86400

    def _get_key(self, idempotency_key: str, method_name: str) -> str:
        raw = f"{method_name}:{idempotency_key}"
        return hashlib.sha256(raw.encode()).hexdigest()

    def get_cached_result(
        self, idempotency_key: str, method_name: str
    ) -> dict[str, Any] | None:
        storage_key = self._get_key(idempotency_key, method_name)
        entry = self._storage.get(storage_key)
        if entry is None:
            return None
        result_json, timestamp = entry
        if (datetime.now() - timestamp).total_seconds() > self._ttl_seconds:
            del self._storage[storage_key]
            return None
        try:
            return json.loads(result_json)
        except json.JSONDecodeError:
            return None

    def cache_result(
        self, idempotency_key: str, method_name: str, result: dict[str, Any]
    ) -> None:
        storage_key = self._get_key(idempotency_key, method_name)
        try:
            result_json = json.dumps(result, default=str)
        except TypeError:
            result_json = json.dumps({"result": str(result)}, default=str)
        self._storage[storage_key] = (result_json, datetime.now())


_idempotency_manager = IdempotencyManager()


router = APIRouter()


# ============================================================================
# PYDANTIC MODELS
# ============================================================================


class SupplierStatusEnum(str, Enum):
    ACTIVE = "active"
    INACTIVE = "inactive"
    BLOCKED = "blocked"
    SUSPENDED = "suspended"
    BLACKLISTED = "blacklisted"
    DRAFT = "draft"


class SupplierTypeEnum(str, Enum):
    INDIVIDUAL = "individual"
    COMPANY = "company"
    GOVERNMENT = "government"
    NON_PROFIT = "non_profit"


class WithholdingCategoryEnum(str, Enum):
    NONE = "none"
    PPH23 = "pph23"
    PPH26 = "pph26"
    BOTH = "both"


class CreateSupplierRequest(BaseModel):
    supplier_code: str = Field(..., min_length=1, max_length=30)
    name: str = Field(..., min_length=1, max_length=200)
    company_name: str | None = Field(None, max_length=200)
    supplier_type: SupplierTypeEnum = SupplierTypeEnum.COMPANY
    npwp: str | None = Field(None, max_length=20)
    tax_name: str | None = Field(None, max_length=200)
    address: str | None = None
    city: str | None = Field(None, max_length=100)
    province: str | None = Field(None, max_length=100)
    postal_code: str | None = Field(None, max_length=20)
    country: str = Field("ID", max_length=2)
    phone: str | None = Field(None, max_length=20)
    mobile: str | None = Field(None, max_length=20)
    email: str | None = Field(None, max_length=200)
    website: str | None = Field(None, max_length=200)
    contact_person: str | None = Field(None, max_length=100)
    payment_terms_days: int = Field(30, ge=0, le=365)
    credit_limit: Decimal = Field(Decimal("0"), ge=0)
    opening_balance: Decimal = Field(Decimal("0"), ge=0)
    opening_balance_date: date | None = None
    bank_name: str | None = Field(None, max_length=100)
    bank_account_number: str | None = Field(None, max_length=50)
    bank_account_name: str | None = Field(None, max_length=100)
    withholding_category: WithholdingCategoryEnum = WithholdingCategoryEnum.NONE
    remarks: str | None = None

    model_config = ConfigDict(use_enum_values=True)


class UpdateSupplierRequest(BaseModel):
    name: str | None = Field(None, max_length=200)
    company_name: str | None = Field(None, max_length=200)
    supplier_type: SupplierTypeEnum | None = None
    npwp: str | None = Field(None, max_length=20)
    tax_name: str | None = Field(None, max_length=200)
    address: str | None = None
    city: str | None = Field(None, max_length=100)
    province: str | None = Field(None, max_length=100)
    postal_code: str | None = Field(None, max_length=20)
    country: str | None = Field(None, max_length=2)
    phone: str | None = Field(None, max_length=20)
    mobile: str | None = Field(None, max_length=20)
    email: str | None = Field(None, max_length=200)
    website: str | None = Field(None, max_length=200)
    contact_person: str | None = Field(None, max_length=100)
    payment_terms_days: int | None = Field(None, ge=0, le=365)
    credit_limit: Decimal | None = Field(None, ge=0)
    opening_balance: Decimal | None = Field(None, ge=0)
    opening_balance_date: date | None = None
    bank_name: str | None = Field(None, max_length=100)
    bank_account_number: str | None = Field(None, max_length=50)
    bank_account_name: str | None = Field(None, max_length=100)
    withholding_category: WithholdingCategoryEnum | None = None
    remarks: str | None = None
    is_active: bool | None = None
    status: SupplierStatusEnum | None = None

    model_config = ConfigDict(use_enum_values=True)


class SupplierResponseModel(BaseModel):
    id: UUID
    legal_entity_id: UUID
    supplier_code: str
    name: str
    company_name: str | None
    supplier_type: str
    npwp: str | None
    tax_name: str | None
    address: str | None
    city: str | None
    province: str | None
    postal_code: str | None
    country: str
    phone: str | None
    mobile: str | None
    email: str | None
    website: str | None
    contact_person: str | None
    payment_terms_days: int
    credit_limit: Decimal
    opening_balance: Decimal
    opening_balance_date: date | None
    bank_name: str | None
    bank_account_number: str | None
    bank_account_name: str | None
    withholding_category: str
    remarks: str | None
    is_active: bool
    status: str
    created_at: datetime
    updated_at: datetime
    created_by: UUID | None
    version: int


class UpdateWithholdingCategoryRequest(BaseModel):
    withholding_category: WithholdingCategoryEnum


class SupplierListResponse(BaseModel):
    items: list[SupplierResponseModel]
    total: int


class SupplierBalanceResponse(BaseModel):
    supplier_id: UUID
    outstanding_balance: Decimal


class AttachmentRequest(BaseModel):
    document_type: str = Field(
        "other", max_length=30,
        description="npwp|siup|ktp|kontrak|foto|other",
    )
    file_name: str = Field(..., min_length=1, max_length=255)
    file_path: str = Field(..., min_length=1)
    file_size_bytes: int | None = None
    mime_type: str | None = Field(None, max_length=100)
    notes: str | None = None


class NoteRequest(BaseModel):
    note: str = Field(..., min_length=1)

class SupplierAddressRequest(BaseModel):
    address_type: str = Field("other", max_length=20,
        description="billing|shipping|warehouse|other")
    label: str | None = Field(None, max_length=100)
    address_line: str = Field(..., min_length=1)
    city: str | None = Field(None, max_length=100)
    province: str | None = Field(None, max_length=100)
    district: str | None = Field(None, max_length=100)
    postal_code: str | None = Field(None, max_length=20)
    country: str = Field("ID", max_length=2)
    latitude: Decimal | None = None
    longitude: Decimal | None = None
    is_primary: bool = False


class SupplierContactRequest(BaseModel):
    name: str = Field(..., min_length=1, max_length=150)
    position: str | None = Field(None, max_length=100)
    phone: str | None = Field(None, max_length=20)
    mobile: str | None = Field(None, max_length=20)
    email: str | None = Field(None, max_length=200)
    whatsapp: str | None = Field(None, max_length=20)
    is_primary: bool = False


class SupplierTagRequest(BaseModel):
    tag: str = Field(..., min_length=1, max_length=50)


class UpdateCreditLimitRequest(BaseModel):
    new_limit: Decimal = Field(..., ge=0)
    reason: str | None = None


# ============================================================================
# HELPERS
# ============================================================================


def get_correlation_id(request: Request) -> str:
    corr_id = request.headers.get("X-Correlation-ID")
    if not corr_id:
        from uuid import uuid4
        corr_id = str(uuid4())
    return corr_id


def to_supplier_response(supplier: Supplier) -> SupplierResponseModel:
    return SupplierResponseModel(
        id=supplier.id,
        legal_entity_id=supplier.legal_entity_id,
        supplier_code=supplier.supplier_code,
        name=supplier.name,
        company_name=supplier.company_name,
        supplier_type=supplier.supplier_type,
        npwp=supplier.npwp,
        tax_name=supplier.tax_name,
        address=supplier.address,
        city=supplier.city,
        province=supplier.province,
        postal_code=supplier.postal_code,
        country=supplier.country,
        phone=supplier.phone,
        mobile=supplier.mobile,
        email=supplier.email,
        website=supplier.website,
        contact_person=supplier.contact_person,
        payment_terms_days=supplier.payment_terms_days,
        credit_limit=supplier.credit_limit,
        opening_balance=supplier.opening_balance,
        opening_balance_date=supplier.opening_balance_date,
        bank_name=supplier.bank_name,
        bank_account_number=supplier.bank_account_number,
        bank_account_name=supplier.bank_account_name,
        withholding_category=supplier.withholding_category,
        remarks=supplier.remarks,
        is_active=supplier.is_active,
        status=supplier.status.value
        if hasattr(supplier.status, "value")
        else supplier.status,
        created_at=supplier.created_at,
        updated_at=supplier.updated_at,
        created_by=supplier.created_by,
        version=supplier.version,
    )


# ============================================================================
# EXPORT COLUMNS (didefinisikan sekali, dipakai CSV & Excel)
# ============================================================================

_EXPORT_COLUMNS: list[tuple[str, str]] = [
    ("supplier_code", "Kode"),
    ("name", "Nama Supplier"),
    ("company_name", "Perusahaan"),
    ("supplier_type", "Tipe"),
    ("npwp", "NPWP"),
    ("tax_name", "Nama Wajib Pajak"),
    ("withholding_category", "Kategori PPh"),
    ("address", "Alamat"),
    ("city", "Kota"),
    ("province", "Provinsi"),
    ("postal_code", "Kode Pos"),
    ("country", "Negara"),
    ("phone", "Telepon"),
    ("mobile", "HP"),
    ("email", "Email"),
    ("website", "Website"),
    ("contact_person", "Contact Person"),
    ("payment_terms_days", "Termin (hari)"),
    ("credit_limit", "Limit Kredit"),
    ("opening_balance", "Saldo Awal"),
    ("bank_name", "Nama Bank"),
    ("bank_account_number", "No. Rekening"),
    ("bank_account_name", "Nama Pemilik Rekening"),
    ("status", "Status"),
    ("is_active", "Aktif"),
    ("remarks", "Catatan"),
    ("created_at", "Dibuat"),
]
_EXPORT_NUMERIC = {
    "credit_limit",
    "opening_balance",
    "payment_terms_days",
}


def _export_cell(field: str, value: Any) -> Any:
    if value is None:
        return ""
    if field in _EXPORT_NUMERIC:
        try:
            return float(value)
        except (TypeError, ValueError):
            return 0.0
    if field == "is_active":
        return "Ya" if value else "Tidak"
    if isinstance(value, datetime):
        return value.strftime("%Y-%m-%d %H:%M")
    return str(value)


# ============================================================================
# SUPPLIER CRUD ENDPOINTS
# ============================================================================


@router.post(
    "/suppliers",
    response_model=SupplierResponseModel,
    status_code=status.HTTP_201_CREATED,
    summary="Create a new supplier",
)
async def create_supplier(
    request: Request,
    payload: CreateSupplierRequest,
    legal_entity_id: UUID = Depends(get_current_legal_entity),
    idempotency_key: str | None = Header(None, alias="Idempotency-Key"),
    user: TokenPayload = Depends(get_current_user),
    service: SupplierService = Depends(get_service(SupplierService)),
) -> SupplierResponseModel:
    method_name = "create_supplier"
    if idempotency_key:
        cached = _idempotency_manager.get_cached_result(
            idempotency_key, method_name
        )
        if cached is not None:
            return SupplierResponseModel(**cached)

    try:
        correlation_id = get_correlation_id(request)
        result = await service.create_supplier(
            legal_entity_id=legal_entity_id,
            supplier_code=payload.supplier_code,
            name=payload.name,
            company_name=payload.company_name,
            supplier_type=payload.supplier_type,
            npwp=payload.npwp,
            tax_name=payload.tax_name,
            address=payload.address,
            city=payload.city,
            province=payload.province,
            postal_code=payload.postal_code,
            country=payload.country,
            phone=payload.phone,
            mobile=payload.mobile,
            email=payload.email,
            website=payload.website,
            contact_person=payload.contact_person,
            payment_terms_days=payload.payment_terms_days,
            credit_limit=payload.credit_limit,
            opening_balance=payload.opening_balance,
            opening_balance_date=payload.opening_balance_date,
            bank_name=payload.bank_name,
            bank_account_number=payload.bank_account_number,
            bank_account_name=payload.bank_account_name,
            withholding_category=payload.withholding_category,
            remarks=payload.remarks,
            created_by=user.user_id,
            correlation_id=correlation_id,
        )
        response = to_supplier_response(result)

        if idempotency_key:
            _idempotency_manager.cache_result(
                idempotency_key, method_name, response.model_dump()
            )

        return response
    except SupplierServiceError as e:
        logger.warning(f"Supplier service error: {e}")
        raise HTTPException(
            status_code=status.HTTP_400_BAD_REQUEST, detail=str(e)
        )
    except Exception as e:
        logger.error(f"Error creating supplier: {e}", exc_info=True)
        raise HTTPException(
            status_code=status.HTTP_500_INTERNAL_SERVER_ERROR, detail=str(e)
        )


# ============================================================================
# EXPORT / IMPORT
# ============================================================================


@router.get(
    "/export",
    summary="Export daftar supplier (csv | excel | json)",
)
async def export_suppliers(
    format: str = Query("csv", pattern="^(csv|excel|json)$"),
    legal_entity_id: UUID = Depends(get_current_legal_entity),
    user: TokenPayload = Depends(get_current_user),
    service: SupplierService = Depends(get_service(SupplierService)),
) -> Response:
    """
    Export suppliers ke 3 format:
    - csv   : text/csv dengan BOM (kompatibel Excel)
    - excel : file .xlsx asli (openpyxl)
    - json  : application/json
    """
    try:
        results, _total = await service.list_suppliers(
            legal_entity_id=legal_entity_id,
            limit=100000,
            offset=0,
        )
        headers_row = [label for _f, label in _EXPORT_COLUMNS]
        rows = [
            [_export_cell(f, getattr(s, f, None)) for f, _l in _EXPORT_COLUMNS]
            for s in results
        ]
        stamp = datetime.now().strftime("%Y%m%d_%H%M")

        if format == "json":
            payload = [dict(zip(headers_row, r)) for r in rows]
            return Response(
                content=json.dumps(payload, ensure_ascii=False, indent=2),
                media_type="application/json",
                headers={
                    "Content-Disposition": (
                        f'attachment; filename="suppliers_{stamp}.json"'
                    )
                },
            )

        if format == "csv":
            buf = io.StringIO()
            writer = csv.writer(buf)
            writer.writerow(headers_row)
            writer.writerows(rows)
            # BOM supaya Excel membaca UTF-8 dengan benar.
            return Response(
                content=("\ufeff" + buf.getvalue()).encode("utf-8"),
                media_type="text/csv; charset=utf-8",
                headers={
                    "Content-Disposition": (
                        f'attachment; filename="suppliers_{stamp}.csv"'
                    )
                },
            )

        # format == "excel" -> .xlsx asli via openpyxl
        try:
            from openpyxl import Workbook
            from openpyxl.styles import Alignment, Font, PatternFill
            from openpyxl.utils import get_column_letter
        except ImportError as exc:
            raise HTTPException(
                status_code=status.HTTP_501_NOT_IMPLEMENTED,
                detail=(
                    "Export Excel membutuhkan paket 'openpyxl' "
                    "(pip install openpyxl). Gunakan format csv."
                ),
            ) from exc
        wb = Workbook()
        ws = wb.active
        ws.title = "Pemasok"
        ws.append(headers_row)
        for cell in ws[1]:
            cell.font = Font(name="Arial", bold=True, color="FFFFFF")
            cell.fill = PatternFill("solid", fgColor="0891B2")
            cell.alignment = Alignment(vertical="center", wrap_text=True)
        for r in rows:
            ws.append(r)
        for i, (field, label) in enumerate(_EXPORT_COLUMNS, start=1):
            col = get_column_letter(i)
            longest = max(
                [len(label)] + [len(str(r[i - 1])) for r in rows[:500]]
            )
            ws.column_dimensions[col].width = min(max(longest + 2, 10), 45)
            if field in ("credit_limit", "opening_balance"):
                for cell in ws[col][1:]:
                    cell.number_format = "#,##0.00"
        ws.freeze_panes = "A2"
        ws.auto_filter.ref = ws.dimensions
        out = io.BytesIO()
        wb.save(out)
        return Response(
            content=out.getvalue(),
            media_type=(
                "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
            ),
            headers={
                "Content-Disposition": (
                    f'attachment; filename="suppliers_{stamp}.xlsx"'
                )
            },
        )
    except HTTPException:
        raise
    except Exception as e:
        logger.error(f"Error exporting suppliers: {e}", exc_info=True)
        raise HTTPException(
            status_code=status.HTTP_500_INTERNAL_SERVER_ERROR, detail=str(e)
        )


@router.post("/import", summary="Import suppliers from CSV")
async def import_suppliers(
    request: Request,
    rows: list[CreateSupplierRequest] = Body(...),
    legal_entity_id: UUID = Depends(get_current_legal_entity),
    user: TokenPayload = Depends(get_current_user),
    service: SupplierService = Depends(get_service(SupplierService)),
) -> dict[str, Any]:
    correlation_id = get_correlation_id(request)
    created = 0
    errors: list[dict[str, Any]] = []
    for idx, row in enumerate(rows):
        try:
            await service.create_supplier(
                legal_entity_id=legal_entity_id,
                supplier_code=row.supplier_code,
                name=row.name,
                company_name=row.company_name,
                supplier_type=row.supplier_type,
                npwp=row.npwp,
                tax_name=row.tax_name,
                address=row.address,
                city=row.city,
                province=row.province,
                postal_code=row.postal_code,
                country=row.country,
                phone=row.phone,
                mobile=row.mobile,
                email=row.email,
                website=row.website,
                contact_person=row.contact_person,
                payment_terms_days=row.payment_terms_days,
                credit_limit=row.credit_limit,
                opening_balance=row.opening_balance,
                opening_balance_date=row.opening_balance_date,
                bank_name=row.bank_name,
                bank_account_number=row.bank_account_number,
                bank_account_name=row.bank_account_name,
                withholding_category=row.withholding_category,
                remarks=row.remarks,
                created_by=user.user_id,
                correlation_id=correlation_id,
            )
            created += 1
        except SupplierServiceError as e:
            errors.append(
                {"row": idx, "supplier_code": row.supplier_code, "error": str(e)}
            )
        except Exception as e:
            logger.error(
                f"Unexpected error importing row {idx}: {e}", exc_info=True
            )
            errors.append(
                {"row": idx, "supplier_code": row.supplier_code, "error": str(e)}
            )

    return {"created": created, "failed": len(errors), "errors": errors}


@router.get("/next-code", summary="Get next auto-generated supplier code")
async def get_next_supplier_code(
    prefix: str = Query("SUP"),
    legal_entity_id: UUID = Depends(get_current_legal_entity),
    user: TokenPayload = Depends(get_current_user),
    service: SupplierService = Depends(get_service(SupplierService)),
) -> dict[str, str]:
    try:
        next_code = await service.get_next_code(legal_entity_id, prefix)
        return {"supplier_code": next_code}
    except Exception as e:
        logger.error(f"Error getting next supplier code: {e}", exc_info=True)
        raise HTTPException(
            status_code=status.HTTP_500_INTERNAL_SERVER_ERROR, detail=str(e)
        )


@router.get("/stats", summary="Get supplier statistics")
async def get_supplier_stats(
    legal_entity_id: UUID = Depends(get_current_legal_entity),
    user: TokenPayload = Depends(get_current_user),
    service: SupplierService = Depends(get_service(SupplierService)),
) -> dict[str, Any]:
    try:
        return await service.get_statistics(legal_entity_id)
    except Exception as e:
        logger.error(f"Error getting supplier stats: {e}", exc_info=True)
        raise HTTPException(
            status_code=status.HTTP_500_INTERNAL_SERVER_ERROR, detail=str(e)
        )


@router.get(
    "/suppliers/{supplier_id}",
    response_model=SupplierResponseModel,
    summary="Get supplier by ID",
)
async def get_supplier(
    supplier_id: UUID,
    legal_entity_id: UUID = Depends(get_current_legal_entity),
    user: TokenPayload = Depends(get_current_user),
    service: SupplierService = Depends(get_service(SupplierService)),
) -> SupplierResponseModel:
    try:
        result = await service.get_supplier(supplier_id, legal_entity_id)
        if not result:
            raise HTTPException(
                status_code=status.HTTP_404_NOT_FOUND, detail="Supplier not found"
            )
        return to_supplier_response(result)
    except HTTPException:
        raise
    except Exception as e:
        logger.error(f"Error getting supplier {supplier_id}: {e}", exc_info=True)
        raise HTTPException(
            status_code=status.HTTP_500_INTERNAL_SERVER_ERROR, detail=str(e)
        )


@router.get(
    "/suppliers", response_model=SupplierListResponse, summary="List suppliers"
)
async def list_suppliers(
    legal_entity_id: UUID = Depends(get_current_legal_entity),
    search: str | None = Query(None),
    city: str | None = Query(None),
    is_active: bool | None = Query(None),
    status_filter: SupplierStatusEnum | None = Query(None, alias="status"),
    page: int = Query(1, ge=1),
    page_size: int | None = Query(None, ge=1, le=1000),
    limit: int = Query(50, ge=1, le=1000),
    offset: int = Query(0, ge=0),
    user: TokenPayload = Depends(get_current_user),
    service: SupplierService = Depends(get_service(SupplierService)),
) -> SupplierListResponse:
    try:
        effective_limit = page_size or limit
        effective_offset = (page - 1) * effective_limit if page_size else offset

        results, total = await service.list_suppliers(
            legal_entity_id=legal_entity_id,
            search=search,
            city=city,
            is_active=is_active,
            status=status_filter.value if status_filter else None,
            limit=effective_limit,
            offset=effective_offset,
        )
        return SupplierListResponse(
            items=[to_supplier_response(s) for s in results],
            total=total,
        )
    except Exception as e:
        logger.error(f"Error listing suppliers: {e}", exc_info=True)
        raise HTTPException(
            status_code=status.HTTP_500_INTERNAL_SERVER_ERROR, detail=str(e)
        )


@router.patch(
    "/suppliers/{supplier_id}",
    response_model=SupplierResponseModel,
    summary="Update supplier",
)
async def update_supplier(
    request: Request,
    supplier_id: UUID,
    payload: UpdateSupplierRequest,
    legal_entity_id: UUID = Depends(get_current_legal_entity),
    idempotency_key: str | None = Header(None, alias="Idempotency-Key"),
    user: TokenPayload = Depends(get_current_user),
    service: SupplierService = Depends(get_service(SupplierService)),
) -> SupplierResponseModel:
    method_name = "update_supplier"
    if idempotency_key:
        cached = _idempotency_manager.get_cached_result(
            idempotency_key, method_name
        )
        if cached is not None:
            return SupplierResponseModel(**cached)

    try:
        correlation_id = get_correlation_id(request)
        result = await service.update_supplier(
            supplier_id=supplier_id,
            legal_entity_id=legal_entity_id,
            name=payload.name,
            company_name=payload.company_name,
            supplier_type=payload.supplier_type,
            npwp=payload.npwp,
            tax_name=payload.tax_name,
            address=payload.address,
            city=payload.city,
            province=payload.province,
            postal_code=payload.postal_code,
            country=payload.country,
            phone=payload.phone,
            mobile=payload.mobile,
            email=payload.email,
            website=payload.website,
            contact_person=payload.contact_person,
            payment_terms_days=payload.payment_terms_days,
            credit_limit=payload.credit_limit,
            opening_balance=payload.opening_balance,
            opening_balance_date=payload.opening_balance_date,
            bank_name=payload.bank_name,
            bank_account_number=payload.bank_account_number,
            bank_account_name=payload.bank_account_name,
            withholding_category=payload.withholding_category,
            remarks=payload.remarks,
            is_active=payload.is_active,
            status=payload.status if payload.status else None,
            updated_by=user.user_id,
            correlation_id=correlation_id,
        )
        response = to_supplier_response(result)

        if idempotency_key:
            _idempotency_manager.cache_result(
                idempotency_key, method_name, response.model_dump()
            )

        return response
    except SupplierNotFoundError:
        raise HTTPException(
            status_code=status.HTTP_404_NOT_FOUND, detail="Supplier not found"
        )
    except SupplierServiceError as e:
        logger.warning(f"Supplier service error: {e}")
        raise HTTPException(
            status_code=status.HTTP_400_BAD_REQUEST, detail=str(e)
        )
    except Exception as e:
        logger.error(
            f"Error updating supplier {supplier_id}: {e}", exc_info=True
        )
        raise HTTPException(
            status_code=status.HTTP_500_INTERNAL_SERVER_ERROR, detail=str(e)
        )


# ============================================================================
# STATUS MANAGEMENT
# ============================================================================


@router.post(
    "/suppliers/{supplier_id}/deactivate",
    status_code=status.HTTP_204_NO_CONTENT,
    summary="Deactivate supplier",
)
async def deactivate_supplier(
    request: Request,
    supplier_id: UUID,
    legal_entity_id: UUID = Depends(get_current_legal_entity),
    idempotency_key: str | None = Header(None, alias="Idempotency-Key"),
    user: TokenPayload = Depends(get_current_user),
    service: SupplierService = Depends(get_service(SupplierService)),
) -> Response:
    try:
        correlation_id = get_correlation_id(request)
        supplier = await service.get_supplier(supplier_id, legal_entity_id)
        if not supplier:
            raise HTTPException(
                status_code=status.HTTP_404_NOT_FOUND, detail="Supplier not found"
            )
        await service.update_supplier(
            supplier_id=supplier_id,
            legal_entity_id=legal_entity_id,
            is_active=False,
            status=SupplierStatusEnum.INACTIVE.value,
            updated_by=user.user_id,
            correlation_id=correlation_id,
        )
        return Response(status_code=status.HTTP_204_NO_CONTENT)
    except HTTPException:
        raise
    except Exception as e:
        logger.error(
            f"Error deactivating supplier {supplier_id}: {e}", exc_info=True
        )
        raise HTTPException(
            status_code=status.HTTP_500_INTERNAL_SERVER_ERROR, detail=str(e)
        )


@router.post(
    "/suppliers/{supplier_id}/activate",
    status_code=status.HTTP_204_NO_CONTENT,
    summary="Activate supplier",
)
async def activate_supplier(
    request: Request,
    supplier_id: UUID,
    legal_entity_id: UUID = Depends(get_current_legal_entity),
    idempotency_key: str | None = Header(None, alias="Idempotency-Key"),
    user: TokenPayload = Depends(get_current_user),
    service: SupplierService = Depends(get_service(SupplierService)),
) -> Response:
    try:
        correlation_id = get_correlation_id(request)
        supplier = await service.get_supplier(supplier_id, legal_entity_id)
        if not supplier:
            raise HTTPException(
                status_code=status.HTTP_404_NOT_FOUND, detail="Supplier not found"
            )
        await service.update_supplier(
            supplier_id=supplier_id,
            legal_entity_id=legal_entity_id,
            is_active=True,
            status=SupplierStatusEnum.ACTIVE.value,
            updated_by=user.user_id,
            correlation_id=correlation_id,
        )
        return Response(status_code=status.HTTP_204_NO_CONTENT)
    except HTTPException:
        raise
    except Exception as e:
        logger.error(
            f"Error activating supplier {supplier_id}: {e}", exc_info=True
        )
        raise HTTPException(
            status_code=status.HTTP_500_INTERNAL_SERVER_ERROR, detail=str(e)
        )


@router.delete(
    "/suppliers/{supplier_id}",
    status_code=status.HTTP_204_NO_CONTENT,
    summary="Delete (soft-delete) a supplier",
)
async def delete_supplier(
    request: Request,
    supplier_id: UUID,
    legal_entity_id: UUID = Depends(get_current_legal_entity),
    idempotency_key: str | None = Header(None, alias="Idempotency-Key"),
    user: TokenPayload = Depends(get_current_user),
    service: SupplierService = Depends(get_service(SupplierService)),
) -> Response:
    try:
        correlation_id = get_correlation_id(request)
        await service.delete_supplier(
            supplier_id=supplier_id,
            legal_entity_id=legal_entity_id,
            deleted_by=user.user_id,
        )
        return Response(status_code=status.HTTP_204_NO_CONTENT)
    except SupplierNotFoundError:
        raise HTTPException(
            status_code=status.HTTP_404_NOT_FOUND, detail="Supplier not found"
        )
    except SupplierHasTransactionsError as e:
        raise HTTPException(status_code=status.HTTP_409_CONFLICT, detail=str(e))
    except Exception as e:
        logger.error(
            f"Error deleting supplier {supplier_id}: {e}", exc_info=True
        )
        raise HTTPException(
            status_code=status.HTTP_500_INTERNAL_SERVER_ERROR, detail=str(e)
        )


@router.post(
    "/suppliers/{supplier_id}/status",
    response_model=SupplierResponseModel,
    summary="Change supplier status",
)
async def change_supplier_status(
    request: Request,
    supplier_id: UUID,
    new_status: SupplierStatusEnum = Body(...),
    legal_entity_id: UUID = Depends(get_current_legal_entity),
    idempotency_key: str | None = Header(None, alias="Idempotency-Key"),
    user: TokenPayload = Depends(get_current_user),
    service: SupplierService = Depends(get_service(SupplierService)),
) -> SupplierResponseModel:
    method_name = "change_supplier_status"
    if idempotency_key:
        cached = _idempotency_manager.get_cached_result(
            idempotency_key, method_name
        )
        if cached is not None:
            return SupplierResponseModel(**cached)

    try:
        correlation_id = get_correlation_id(request)
        result = await service.update_supplier(
            supplier_id=supplier_id,
            legal_entity_id=legal_entity_id,
            status=new_status.value,
            updated_by=user.user_id,
            correlation_id=correlation_id,
        )
        response = to_supplier_response(result)
        if idempotency_key:
            _idempotency_manager.cache_result(
                idempotency_key, method_name, response.model_dump()
            )
        return response
    except SupplierNotFoundError:
        raise HTTPException(
            status_code=status.HTTP_404_NOT_FOUND, detail="Supplier not found"
        )
    except Exception as e:
        logger.error(
            f"Error changing supplier status {supplier_id}: {e}", exc_info=True
        )
        raise HTTPException(
            status_code=status.HTTP_500_INTERNAL_SERVER_ERROR, detail=str(e)
        )


# ============================================================================
# WITHHOLDING CATEGORY
# ============================================================================


@router.post(
    "/suppliers/{supplier_id}/withholding-category",
    response_model=SupplierResponseModel,
    summary="Update withholding category",
)
async def update_withholding_category(
    request: Request,
    supplier_id: UUID,
    payload: UpdateWithholdingCategoryRequest,
    legal_entity_id: UUID = Depends(get_current_legal_entity),
    idempotency_key: str | None = Header(None, alias="Idempotency-Key"),
    user: TokenPayload = Depends(get_current_user),
    service: SupplierService = Depends(get_service(SupplierService)),
) -> SupplierResponseModel:
    try:
        correlation_id = get_correlation_id(request)
        result = await service.update_withholding_category(
            supplier_id=supplier_id,
            legal_entity_id=legal_entity_id,
            withholding_category=payload.withholding_category.value,
            updated_by=user.user_id,
            correlation_id=correlation_id,
        )
        return to_supplier_response(result)
    except SupplierNotFoundError:
        raise HTTPException(
            status_code=status.HTTP_404_NOT_FOUND, detail="Supplier not found"
        )
    except Exception as e:
        logger.error(
            f"Error updating withholding category for supplier "
            f"{supplier_id}: {e}",
            exc_info=True,
        )
        raise HTTPException(
            status_code=status.HTTP_500_INTERNAL_SERVER_ERROR, detail=str(e)
        )


# ============================================================================
# BALANCE
# ============================================================================


@router.get(
    "/suppliers/{supplier_id}/balance",
    response_model=SupplierBalanceResponse,
    summary="Get supplier outstanding AP balance",
)
async def get_supplier_balance(
    supplier_id: UUID,
    legal_entity_id: UUID = Depends(get_current_legal_entity),
    user: TokenPayload = Depends(get_current_user),
    service: SupplierService = Depends(get_service(SupplierService)),
) -> SupplierBalanceResponse:
    try:
        supplier = await service.get_supplier(supplier_id, legal_entity_id)
        if not supplier:
            raise HTTPException(
                status_code=status.HTTP_404_NOT_FOUND, detail="Supplier not found"
            )
        balance = await service.get_outstanding_balance(supplier_id)
        return SupplierBalanceResponse(
            supplier_id=supplier_id, outstanding_balance=balance
        )
    except HTTPException:
        raise
    except Exception as e:
        logger.error(
            f"Error getting balance for supplier {supplier_id}: {e}",
            exc_info=True,
        )
        raise HTTPException(
            status_code=status.HTTP_500_INTERNAL_SERVER_ERROR, detail=str(e)
        )


# ============================================================================
# AUDIT HISTORY
# ============================================================================


@router.get(
    "/suppliers/{supplier_id}/history",
    summary="Get supplier change history (audit trail)",
)
async def get_supplier_history(
    supplier_id: UUID,
    legal_entity_id: UUID = Depends(get_current_legal_entity),
    user: TokenPayload = Depends(get_current_user),
    service: SupplierService = Depends(get_service(SupplierService)),
) -> list[dict[str, Any]]:
    supplier = await service.get_supplier(supplier_id, legal_entity_id)
    if not supplier:
        raise HTTPException(
            status_code=status.HTTP_404_NOT_FOUND, detail="Supplier not found"
        )
    return [
        entry
        for entry in service.get_audit_trail()
        if entry.get("details", {}).get("supplier_id") == str(supplier_id)
    ]


# ============================================================================
# SUPPLIER ATTACHMENTS
# ============================================================================


@router.get(
    "/suppliers/{supplier_id}/attachments",
    response_model=list[dict],
    summary="List supplier attachments",
)
async def list_supplier_attachments(
    supplier_id: UUID,
    legal_entity_id: UUID = Depends(get_current_legal_entity),
    user: TokenPayload = Depends(get_current_user),
    service: SupplierService = Depends(get_service(SupplierService)),
) -> list[dict]:
    try:
        return await service.list_attachments(supplier_id, legal_entity_id)
    except SupplierNotFoundError:
        raise HTTPException(
            status_code=status.HTTP_404_NOT_FOUND, detail="Supplier not found"
        )
    except Exception as e:
        logger.error(f"Error listing attachments: {e}", exc_info=True)
        raise HTTPException(
            status_code=status.HTTP_500_INTERNAL_SERVER_ERROR, detail=str(e)
        )


@router.post(
    "/suppliers/{supplier_id}/attachments",
    response_model=dict,
    status_code=status.HTTP_201_CREATED,
    summary="Add supplier attachment",
)
async def add_supplier_attachment(
    supplier_id: UUID,
    payload: AttachmentRequest,
    legal_entity_id: UUID = Depends(get_current_legal_entity),
    user: TokenPayload = Depends(get_current_user),
    service: SupplierService = Depends(get_service(SupplierService)),
) -> dict:
    try:
        return await service.add_attachment(
            supplier_id=supplier_id,
            legal_entity_id=legal_entity_id,
            document_type=payload.document_type,
            file_name=payload.file_name,
            file_path=payload.file_path,
            file_size_bytes=payload.file_size_bytes,
            mime_type=payload.mime_type,
            notes=payload.notes,
            created_by=user.user_id,
        )
    except SupplierNotFoundError:
        raise HTTPException(
            status_code=status.HTTP_404_NOT_FOUND, detail="Supplier not found"
        )
    except Exception as e:
        logger.error(f"Error adding attachment: {e}", exc_info=True)
        raise HTTPException(
            status_code=status.HTTP_500_INTERNAL_SERVER_ERROR, detail=str(e)
        )


@router.delete(
    "/suppliers/{supplier_id}/attachments/{attachment_id}",
    status_code=status.HTTP_204_NO_CONTENT,
    response_model=None,
    summary="Delete supplier attachment",
)
async def delete_supplier_attachment(
    supplier_id: UUID,
    attachment_id: UUID,
    user: TokenPayload = Depends(get_current_user),
    service: SupplierService = Depends(get_service(SupplierService)),
) -> None:
    try:
        ok = await service.delete_attachment(
            supplier_id, attachment_id, deleted_by=user.user_id
        )
        if not ok:
            raise HTTPException(
                status_code=status.HTTP_404_NOT_FOUND, detail="Attachment not found"
            )
    except HTTPException:
        raise
    except Exception as e:
        logger.error(f"Error deleting attachment: {e}", exc_info=True)
        raise HTTPException(
            status_code=status.HTTP_500_INTERNAL_SERVER_ERROR, detail=str(e)
        )


# ============================================================================
# SUPPLIER NOTES
# ============================================================================


@router.get(
    "/suppliers/{supplier_id}/notes",
    response_model=list[dict],
    summary="List supplier notes",
)
async def list_supplier_notes(
    supplier_id: UUID,
    legal_entity_id: UUID = Depends(get_current_legal_entity),
    user: TokenPayload = Depends(get_current_user),
    service: SupplierService = Depends(get_service(SupplierService)),
) -> list[dict]:
    try:
        return await service.list_notes(supplier_id, legal_entity_id)
    except SupplierNotFoundError:
        raise HTTPException(
            status_code=status.HTTP_404_NOT_FOUND, detail="Supplier not found"
        )
    except Exception as e:
        logger.error(f"Error listing notes: {e}", exc_info=True)
        raise HTTPException(
            status_code=status.HTTP_500_INTERNAL_SERVER_ERROR, detail=str(e)
        )


@router.post(
    "/suppliers/{supplier_id}/notes",
    response_model=dict,
    status_code=status.HTTP_201_CREATED,
    summary="Add supplier note",
)
async def add_supplier_note(
    supplier_id: UUID,
    payload: NoteRequest,
    legal_entity_id: UUID = Depends(get_current_legal_entity),
    user: TokenPayload = Depends(get_current_user),
    service: SupplierService = Depends(get_service(SupplierService)),
) -> dict:
    try:
        return await service.add_note(
            supplier_id=supplier_id,
            legal_entity_id=legal_entity_id,
            note=payload.note,
            created_by=user.user_id,
        )
    except SupplierNotFoundError:
        raise HTTPException(
            status_code=status.HTTP_404_NOT_FOUND, detail="Supplier not found"
        )
    except SupplierServiceError as e:
        raise HTTPException(
            status_code=status.HTTP_400_BAD_REQUEST, detail=str(e)
        )
    except Exception as e:
        logger.error(f"Error adding note: {e}", exc_info=True)
        raise HTTPException(
            status_code=status.HTTP_500_INTERNAL_SERVER_ERROR, detail=str(e)
        )


@router.delete(
    "/suppliers/{supplier_id}/notes/{note_id}",
    status_code=status.HTTP_204_NO_CONTENT,
    response_model=None,
    summary="Delete supplier note",
)
async def delete_supplier_note(
    supplier_id: UUID,
    note_id: UUID,
    user: TokenPayload = Depends(get_current_user),
    service: SupplierService = Depends(get_service(SupplierService)),
) -> None:
    try:
        ok = await service.delete_note(
            supplier_id, note_id, deleted_by=user.user_id
        )
        if not ok:
            raise HTTPException(
                status_code=status.HTTP_404_NOT_FOUND, detail="Note not found"
            )
    except HTTPException:
        raise
    except Exception as e:
        logger.error(f"Error deleting note: {e}", exc_info=True)
        raise HTTPException(
            status_code=status.HTTP_500_INTERNAL_SERVER_ERROR, detail=str(e)
        )

# ============================================================================
# SUPPLIER ADDRESSES (Tab Alamat)
# ============================================================================


@router.get("/suppliers/{supplier_id}/addresses", response_model=list[dict])
async def list_supplier_addresses(
    supplier_id: UUID,
    legal_entity_id: UUID = Depends(get_current_legal_entity),
    user: TokenPayload = Depends(get_current_user),
    service: SupplierService = Depends(get_service(SupplierService)),
) -> list[dict]:
    try:
        return await service.list_addresses(supplier_id, legal_entity_id)
    except SupplierNotFoundError:
        raise HTTPException(status_code=404, detail="Supplier not found")


@router.post(
    "/suppliers/{supplier_id}/addresses",
    response_model=dict,
    status_code=status.HTTP_201_CREATED,
)
async def add_supplier_address(
    supplier_id: UUID,
    payload: SupplierAddressRequest,
    legal_entity_id: UUID = Depends(get_current_legal_entity),
    user: TokenPayload = Depends(get_current_user),
    service: SupplierService = Depends(get_service(SupplierService)),
) -> dict:
    try:
        return await service.add_address(
            supplier_id, legal_entity_id, **payload.model_dump()
        )
    except SupplierNotFoundError:
        raise HTTPException(status_code=404, detail="Supplier not found")


@router.patch("/suppliers/{supplier_id}/addresses/{address_id}", response_model=dict)
async def update_supplier_address(
    supplier_id: UUID,
    address_id: UUID,
    payload: SupplierAddressRequest,
    user: TokenPayload = Depends(get_current_user),
    service: SupplierService = Depends(get_service(SupplierService)),
) -> dict:
    try:
        return await service.update_address(
            supplier_id, address_id, **payload.model_dump(exclude_unset=True)
        )
    except Exception as e:
        raise HTTPException(status_code=400, detail=str(e))


@router.delete(
    "/suppliers/{supplier_id}/addresses/{address_id}",
    status_code=status.HTTP_204_NO_CONTENT,
    response_model=None,
)
async def delete_supplier_address(
    supplier_id: UUID,
    address_id: UUID,
    user: TokenPayload = Depends(get_current_user),
    service: SupplierService = Depends(get_service(SupplierService)),
) -> None:
    ok = await service.delete_address(supplier_id, address_id)
    if not ok:
        raise HTTPException(status_code=404, detail="Address not found")


# ============================================================================
# SUPPLIER CONTACTS (Tab Contact Person)
# ============================================================================


@router.get("/suppliers/{supplier_id}/contacts", response_model=list[dict])
async def list_supplier_contacts(
    supplier_id: UUID,
    legal_entity_id: UUID = Depends(get_current_legal_entity),
    user: TokenPayload = Depends(get_current_user),
    service: SupplierService = Depends(get_service(SupplierService)),
) -> list[dict]:
    try:
        return await service.list_contacts(supplier_id, legal_entity_id)
    except SupplierNotFoundError:
        raise HTTPException(status_code=404, detail="Supplier not found")


@router.post(
    "/suppliers/{supplier_id}/contacts",
    response_model=dict,
    status_code=status.HTTP_201_CREATED,
)
async def add_supplier_contact(
    supplier_id: UUID,
    payload: SupplierContactRequest,
    legal_entity_id: UUID = Depends(get_current_legal_entity),
    user: TokenPayload = Depends(get_current_user),
    service: SupplierService = Depends(get_service(SupplierService)),
) -> dict:
    try:
        return await service.add_contact(
            supplier_id, legal_entity_id, **payload.model_dump()
        )
    except SupplierNotFoundError:
        raise HTTPException(status_code=404, detail="Supplier not found")


@router.patch("/suppliers/{supplier_id}/contacts/{contact_id}", response_model=dict)
async def update_supplier_contact(
    supplier_id: UUID,
    contact_id: UUID,
    payload: SupplierContactRequest,
    user: TokenPayload = Depends(get_current_user),
    service: SupplierService = Depends(get_service(SupplierService)),
) -> dict:
    try:
        return await service.update_contact(
            supplier_id, contact_id, **payload.model_dump(exclude_unset=True)
        )
    except Exception as e:
        raise HTTPException(status_code=400, detail=str(e))


@router.delete(
    "/suppliers/{supplier_id}/contacts/{contact_id}",
    status_code=status.HTTP_204_NO_CONTENT,
    response_model=None,
)
async def delete_supplier_contact(
    supplier_id: UUID,
    contact_id: UUID,
    user: TokenPayload = Depends(get_current_user),
    service: SupplierService = Depends(get_service(SupplierService)),
) -> None:
    ok = await service.delete_contact(supplier_id, contact_id)
    if not ok:
        raise HTTPException(status_code=404, detail="Contact not found")


# ============================================================================
# SUPPLIER TAGS (Tab Tag)
# ============================================================================


@router.get("/suppliers/{supplier_id}/tags", response_model=list[dict])
async def list_supplier_tags(
    supplier_id: UUID,
    legal_entity_id: UUID = Depends(get_current_legal_entity),
    user: TokenPayload = Depends(get_current_user),
    service: SupplierService = Depends(get_service(SupplierService)),
) -> list[dict]:
    try:
        return await service.list_tags(supplier_id, legal_entity_id)
    except SupplierNotFoundError:
        raise HTTPException(status_code=404, detail="Supplier not found")


@router.post(
    "/suppliers/{supplier_id}/tags",
    response_model=dict,
    status_code=status.HTTP_201_CREATED,
)
async def add_supplier_tag(
    supplier_id: UUID,
    payload: SupplierTagRequest,
    legal_entity_id: UUID = Depends(get_current_legal_entity),
    user: TokenPayload = Depends(get_current_user),
    service: SupplierService = Depends(get_service(SupplierService)),
) -> dict:
    try:
        return await service.add_tag(supplier_id, legal_entity_id, payload.tag)
    except SupplierNotFoundError:
        raise HTTPException(status_code=404, detail="Supplier not found")


@router.delete(
    "/suppliers/{supplier_id}/tags/{tag_id}",
    status_code=status.HTTP_204_NO_CONTENT,
    response_model=None,
)
async def remove_supplier_tag(
    supplier_id: UUID,
    tag_id: UUID,
    user: TokenPayload = Depends(get_current_user),
    service: SupplierService = Depends(get_service(SupplierService)),
) -> None:
    ok = await service.remove_tag(supplier_id, tag_id)
    if not ok:
        raise HTTPException(status_code=404, detail="Tag not found")


# ============================================================================
# CREDIT LIMIT (Tombol Limit Kredit)
# ============================================================================


@router.post(
    "/suppliers/{supplier_id}/credit-limit",
    response_model=SupplierResponseModel,
)
async def update_supplier_credit_limit(
    supplier_id: UUID,
    payload: UpdateCreditLimitRequest,
    legal_entity_id: UUID = Depends(get_current_legal_entity),
    user: TokenPayload = Depends(get_current_user),
    service: SupplierService = Depends(get_service(SupplierService)),
) -> SupplierResponseModel:
    try:
        result = await service.update_credit_limit(
            supplier_id=supplier_id,
            legal_entity_id=legal_entity_id,
            new_limit=payload.new_limit,
            reason=payload.reason,
            updated_by=user.user_id,
        )
        return to_supplier_response(result)
    except SupplierNotFoundError:
        raise HTTPException(status_code=404, detail="Supplier not found")
    except SupplierServiceError as e:
        raise HTTPException(status_code=400, detail=str(e))


@router.get(
    "/suppliers/{supplier_id}/credit-limit/history",
    response_model=list[dict],
)
async def get_supplier_credit_history(
    supplier_id: UUID,
    legal_entity_id: UUID = Depends(get_current_legal_entity),
    limit: int = Query(100, ge=1, le=1000),
    offset: int = Query(0, ge=0),
    user: TokenPayload = Depends(get_current_user),
    service: SupplierService = Depends(get_service(SupplierService)),
) -> list[dict]:
    try:
        return await service.get_credit_history(
            supplier_id, legal_entity_id, limit=limit, offset=offset
        )
    except SupplierNotFoundError:
        raise HTTPException(status_code=404, detail="Supplier not found")
    except Exception as e:
        logger.error(
            f"Error getting credit history for supplier {supplier_id}: {e}",
            exc_info=True,
        )
        raise HTTPException(
            status_code=status.HTTP_500_INTERNAL_SERVER_ERROR, detail=str(e)
        )


# ============================================================================
# BALANCE HISTORY (Tab Riwayat Saldo)
# ============================================================================

@router.get(
    "/suppliers/{supplier_id}/balance/history",
    response_model=list[dict],
)
async def get_supplier_balance_history(
    supplier_id: UUID,
    legal_entity_id: UUID = Depends(get_current_legal_entity),
    limit: int = Query(100, ge=1, le=1000),
    offset: int = Query(0, ge=0),
    user: TokenPayload = Depends(get_current_user),
    service: SupplierService = Depends(get_service(SupplierService)),
) -> list[dict]:
    try:
        return await service.get_balance_history(
            supplier_id, legal_entity_id, limit=limit, offset=offset
        )
    except SupplierNotFoundError:
        raise HTTPException(status_code=404, detail="Supplier not found")
    except Exception as e:
        logger.error(
            f"Error getting balance history for supplier {supplier_id}: {e}",
            exc_info=True,
        )
        raise HTTPException(
            status_code=status.HTTP_500_INTERNAL_SERVER_ERROR, detail=str(e)
        )

    
__all__ = ["router"]
